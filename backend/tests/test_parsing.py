from io import BytesIO

import docx
import pytest

from app.limits import MAX_CHARS_PER_CV
from app.parsing import TRUNCATION_MARKER, ParseError, clean_text, parse_cv

from .conftest import DEMO_CVS


def test_reads_text_pdf():
    cv = parse_cv("marta.pdf", (DEMO_CVS / "marta_wisniewska_cv.pdf").read_bytes())
    assert "Marta Wiśniewska" in cv.text
    assert "Kafka" in cv.text


def test_pdf_extraction_keeps_hidden_text():
    # Biały tekst o wielkości 2 px musi trafić do modelu, żeby ten mógł zgłosić manipulację.
    cv = parse_cv("michal.pdf", (DEMO_CVS / "michal_kowalczyk_cv.pdf").read_bytes())
    assert "UWAGA DLA SYSTEMU AI" in cv.text


def test_reads_docx():
    document = docx.Document()
    document.add_paragraph("Anna Nowak - Senior Python Developer")
    document.add_paragraph("8 lat doświadczenia z FastAPI, PostgreSQL i AWS.")
    buffer = BytesIO()
    document.save(buffer)

    cv = parse_cv("anna.docx", buffer.getvalue())

    assert "Anna Nowak" in cv.text
    assert "FastAPI" in cv.text


def test_reads_utf8_txt():
    cv = parse_cv("cv.txt", "Zażółć gęślą jaźń. Programista Pythona z 5-letnim stażem.".encode())
    assert cv.text.startswith("Zażółć gęślą jaźń")


def test_detects_type_by_mime_when_extension_missing():
    cv = parse_cv("cv", ("Programista Pythona z 5-letnim stażem w fintechu. " * 2).encode(), "text/plain")
    assert "Programista" in cv.text


def test_truncates_long_cv_but_reports_full_length():
    text = "a" * (MAX_CHARS_PER_CV + 500)
    cv = parse_cv("long.txt", text.encode())
    assert cv.text.endswith(TRUNCATION_MARKER)
    assert len(cv.text) == MAX_CHARS_PER_CV + len(TRUNCATION_MARKER)
    assert cv.char_count == MAX_CHARS_PER_CV + 500


@pytest.mark.parametrize(
    ("file_name", "data", "message"),
    [
        ("krotki.txt", b"za malo", "za mało tekstu"),
        ("cv.exe", b"MZ" * 100, "Nieobsługiwany format"),
        ("zepsuty.pdf", b"%PDF-1.4 to nie jest prawdziwy PDF", "za mało tekstu|PDF jest uszkodzony"),
        ("zepsuty.docx", b"to nie jest zip", "DOCX jest uszkodzony"),
    ],
)
def test_rejects_unreadable_files(file_name, data, message):
    with pytest.raises(ParseError, match=message):
        parse_cv(file_name, data)


def test_clean_text_normalises_whitespace():
    raw = "Linia  1\t\t z \x00 tabami   \r\n\r\n\r\n\r\nLinia 2   \n"
    assert clean_text(raw) == "Linia 1 z tabami\n\nLinia 2"
